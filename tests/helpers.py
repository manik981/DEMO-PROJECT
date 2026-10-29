"""Shared test builders. All AI responses here are fixtures - no live API calls."""

from __future__ import annotations

import io
from types import SimpleNamespace
from typing import Any

import docx
import pymupdf as fitz

from src.document_processor import ExtractedDocument, PageText


def make_pdf(pages: list[str]) -> bytes:
    document = fitz.open()
    for text in pages:
        page = document.new_page()
        page.insert_text((72, 72), text, fontsize=11)
    data = document.tobytes()
    document.close()
    return data


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        doc_table = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                doc_table.cell(row_index, column_index).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_document(text: str, page: int | None = 1, filename: str = "sample.txt") -> ExtractedDocument:
    page_texts = [PageText(page_number=page, text=text)]
    return ExtractedDocument(
        filename=filename,
        file_type=filename.rsplit(".", 1)[-1],
        pages=page_texts,
        char_count=len(text),
    )


def empty_analysis_payload() -> dict[str, Any]:
    return {
        "executive_summary": {
            "project_title": "Permit Portal",
            "project_overview": "A city permit portal.",
            "business_objective": "Reduce in-person submissions.",
            "intended_users": ["Residents", "City staff"],
            "project_category": "Web application",
            "overall_complexity": "Medium",
            "complexity_rationale": "Integration and unclear migration scope.",
        },
        "functional_requirements": [],
        "non_functional_requirements": [],
        "technical_requirements": [],
        "missing_information": [],
        "ambiguities": [],
        "risks": [],
        "clarification_questions": [],
        "solution_recommendations": {
            "potential_architecture": "A modular web application with an integration layer.",
            "technology_options": [{"name": "PostgreSQL", "reason": "Relational permit data."}],
            "integration_considerations": ["ERP fee posting"],
            "security_recommendations": ["SSO for staff"],
            "implementation_phases": [
                {"name": "Discovery", "scope": "Confirm requirements", "deliverables": ["Backlog"]}
            ],
            "dependencies_and_assumptions": ["Client provides ERP test access"],
        },
        "complexity_assessment": {
            "level": "Medium",
            "reasons": ["Third-party ERP integration", "Unclear data migration"],
            "complexity_drivers": ["ERP integration"],
            "major_unknowns": ["Data quality"],
        },
        "limitations": ["Only the supplied document was analyzed."],
    }


def populated_analysis_payload(evidence: str = "The portal must support 500 concurrent users") -> dict[str, Any]:
    payload = empty_analysis_payload()
    payload.update(
        {
            "functional_requirements": [
                {
                    "id": "FR-01",
                    "description": "Applicants must submit permit applications online.",
                    "category": "Permitting",
                    "priority": "High",
                    "source_page": 1,
                    "source_section": "Functional requirements",
                    "evidence": evidence,
                    "confidence": "High",
                    "origin": "explicit",
                }
            ],
            "non_functional_requirements": [
                {
                    "id": "NFR-01",
                    "category": "scalability",
                    "description": "Support 500 concurrent users.",
                    "status": "explicit",
                    "priority": "High",
                    "source_page": 1,
                    "evidence": evidence,
                    "confidence": "High",
                }
            ],
            "technical_requirements": [
                {
                    "id": "TR-01",
                    "technology": "Tyler Munis ERP",
                    "category": "integration",
                    "status": "explicitly_specified",
                    "notes": "Fee posting integration.",
                    "source_page": 1,
                    "evidence": evidence,
                    "confidence": "High",
                }
            ],
            "missing_information": [
                {
                    "id": "MI-01",
                    "topic": "Expected transaction volume",
                    "why_it_matters": "Capacity sizing depends on volume.",
                    "priority": "High",
                    "potential_impact": "Risk of under-provisioning.",
                }
            ],
            "ambiguities": [
                {
                    "id": "AM-01",
                    "kind": "vague_requirement",
                    "description": "Reasonable response times is undefined.",
                    "evidence": evidence,
                    "source_page": 1,
                    "status": "possible",
                }
            ],
            "risks": [
                {
                    "id": "RK-01",
                    "description": "ERP integration may be harder than expected.",
                    "category": "integration",
                    "severity": "High",
                    "reason": "No API details were supplied.",
                    "evidence": evidence,
                    "mitigation": "Request ERP API documentation.",
                    "basis": "inferred",
                    "source_page": 1,
                }
            ],
            "clarification_questions": [
                {
                    "id": "CQ-01",
                    "question": "What is the expected application volume?",
                    "why_it_matters": "Sizing and performance targets depend on it.",
                    "priority": "High",
                    "related_to": "MI-01",
                }
            ],
        }
    )
    return payload


def chunk_payload() -> dict[str, Any]:
    return {
        "functional_requirements": [],
        "non_functional_requirements": [],
        "technical_requirements": [],
        "missing_information": [],
        "ambiguities": [],
        "risks": [],
        "chunk_summary": "Section summary.",
    }


class FakeCompletion:
    def __init__(self, content: str, finish_reason: str = "stop", with_usage: bool = True) -> None:
        self.choices = [SimpleNamespace(message=SimpleNamespace(content=content), finish_reason=finish_reason)]
        self.usage = (
            SimpleNamespace(prompt_tokens=100, completion_tokens=50, total_tokens=150)
            if with_usage
            else None
        )


class FakeCompletions:
    def __init__(self, responses: list[Any]) -> None:
        self.responses = list(responses)
        self.calls: list[dict[str, Any]] = []

    def create(self, **kwargs: Any) -> FakeCompletion:
        self.calls.append(kwargs)
        if not self.responses:
            raise AssertionError("FakeGroqClient ran out of scripted responses")
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        if isinstance(response, FakeCompletion):
            return response
        return FakeCompletion(response)


class FakeGroqClient:
    def __init__(self, responses: list[Any]) -> None:
        self.chat = SimpleNamespace(completions=FakeCompletions(responses))

    @property
    def calls(self) -> list[dict[str, Any]]:
        return self.chat.completions.calls
